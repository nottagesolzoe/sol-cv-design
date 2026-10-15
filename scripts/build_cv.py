#!/usr/bin/env python3
"""Build an editable single-column CV from verified JSON; no network operations."""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

VIOLET, RED = '5B20D6', 'C52842'

def hyperlink(paragraph, label, url, size=9.5):
    if not url.startswith(('https://', 'http://', 'mailto:')):
        raise ValueError('Links must use http, https, or mailto')
    el = OxmlElement('w:hyperlink')
    el.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run, props = OxmlElement('w:r'), OxmlElement('w:rPr')
    for tag, value in [('color', VIOLET), ('sz', str(int(size * 2)))]:
        item = OxmlElement('w:' + tag)
        item.set(qn('w:val'), value)
        props.append(item)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label
    run.append(text); el.append(run); paragraph._p.append(el)

def rule(paragraph, edge, color, width):
    borders = OxmlElement('w:pBdr'); line = OxmlElement('w:' + edge)
    for key, value in [('val', 'single'), ('sz', str(width)), ('color', color)]:
        line.set(qn('w:' + key), value)
    borders.append(line); paragraph._p.get_or_add_pPr().append(borders)

def build(data, destination):
    if not data.get('name') or not isinstance(data.get('sections'), list):
        raise ValueError('A name and sections list are required')
    doc = Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.45)
    sec.left_margin = sec.right_margin = Inches(.5)
    style = doc.styles['Normal']; style.font.name = 'Arial'; style.font.size = Pt(9.5)
    style.paragraph_format.space_after = Pt(2); style.paragraph_format.line_spacing = 1.0
    for item in doc.styles:
        for border in list(item.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    title = doc.styles['Title']; title.font.name = 'Arial'; title.font.size = Pt(25)
    title.font.bold = True; title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(1)
    name = doc.add_paragraph(data['name'], 'Title'); rule(name, 'top', VIOLET, 16)
    para = doc.add_paragraph(); run = para.add_run(data.get('headline', ''))
    run.bold = True; run.font.color.rgb = RGBColor.from_string(VIOLET)
    def contacts(para, footer=False):
        text = ([data['name']] if footer else []) + data.get('phones', [])
        run = para.add_run(' | '.join(text)); run.font.size = Pt(7 if footer else 9.5)
        for item in data.get('links', []):
            sep = para.add_run(' | '); sep.font.size = Pt(7 if footer else 9.5)
            hyperlink(para, item['label'], item['url'], 7 if footer else 9.5)
    contacts(doc.add_paragraph())
    for section in data['sections']:
        heading = doc.add_paragraph(); run = heading.add_run(section['title'].upper())
        run.bold = True; run.font.size = Pt(11)
        heading.paragraph_format.space_before = Pt(6)
        heading.paragraph_format.keep_with_next = True
        rule(heading, 'bottom', 'D8D8D3', 4)
        for entry in section.get('entries', []):
            para = doc.add_paragraph(); para.paragraph_format.keep_with_next = bool(entry.get('details') or entry.get('supervisor'))
            para.paragraph_format.tab_stops.add_tab_stop(Inches(7.5), WD_TAB_ALIGNMENT.RIGHT)
            if entry.get('url'):
                hyperlink(para, entry['title'], entry['url'])
            else:
                para.add_run(entry['title']).bold = True
            if entry.get('date'):
                run = para.add_run('\t' + entry['date'])
                run.font.color.rgb = RGBColor.from_string(RED); run.font.size = Pt(8.5)
            if entry.get('supervisor'):
                para = doc.add_paragraph(); para.add_run('Supervisor: ').bold = True
                para.add_run(entry['supervisor'])
            for detail in entry.get('details', []):
                para = doc.add_paragraph('• ' + detail)
                para.paragraph_format.left_indent = Inches(.1)
                para.paragraph_format.first_line_indent = Inches(-.1)
    footer = sec.footer.paragraphs[0]; footer.style = doc.styles['Normal']; contacts(footer, True)
    doc.core_properties.author = data['name']; doc.core_properties.title = data['name'] + ' CV'
    Path(destination).parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    build(json.loads(args.input.read_text(encoding='utf-8')), args.output)
